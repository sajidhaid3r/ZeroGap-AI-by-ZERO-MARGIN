"""
Regression tests covering the behaviors verified during development:
- DOCX end-to-end parsing
- PDF two-column reading-order reconstruction
- Bullet-marker stripping bug fix (leading '-' no longer breaks
  action-verb / weak-opener detection)
"""

import io
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import fitz
from docx import Document

from app.pipeline import run_pipeline
from app.extraction.pdf_extractor import extract_spans, page_dimensions
from app.extraction.column_detector import reconstruct_reading_order, spans_to_lines
from app.scoring.content_quality import analyze_bullets


def _build_sample_docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("jane.doe@gmail.com | (555) 123-4567 | linkedin.com/in/janedoe")
    doc.add_paragraph("Summary")
    doc.add_paragraph("Product manager with 6 years of experience building B2B SaaS products.")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Senior Product Manager, Acme Corp, Jan 2022 - Present")
    doc.add_paragraph("- Led a team of 5 engineers to launch a new billing platform, increasing revenue by 20%")
    doc.add_paragraph("- Responsible for roadmap planning across three product lines")
    doc.add_paragraph("Education")
    doc.add_paragraph("BS Computer Science, State University, 2018")
    doc.add_paragraph("Skills")
    doc.add_paragraph("Product Management, SQL, Roadmapping")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _build_two_column_pdf_bytes() -> bytes:
    doc = fitz.open()
    page = doc.new_page(width=612, height=792)
    left_text = [
        (50, 50, "John Smith"),
        (50, 80, "Experience"),
        (50, 110, "Senior Engineer, TechCo"),
    ]
    right_text = [
        (350, 80, "Skills"),
        (350, 110, "Python, SQL, AWS"),
    ]
    for x, y, text in left_text + right_text:
        page.insert_text((x, y), text, fontsize=11)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_docx_pipeline_extracts_contact_info():
    file_bytes = _build_sample_docx_bytes()
    report = run_pipeline(file_bytes, "sample.docx")

    assert report.parsed_resume.contact.full_name == "Jane Doe"
    assert report.parsed_resume.contact.email == "jane.doe@gmail.com"
    assert report.parsed_resume.contact.phone is not None
    assert report.parsed_resume.contact.linkedin_url is not None


def test_docx_pipeline_finds_all_sections():
    file_bytes = _build_sample_docx_bytes()
    report = run_pipeline(file_bytes, "sample.docx")

    section_names = {s.canonical_name for s in report.parsed_resume.raw_sections}
    assert "summary" in section_names
    assert "experience" in section_names
    assert "education" in section_names
    assert "skills" in section_names


def test_docx_pipeline_produces_overall_score():
    file_bytes = _build_sample_docx_bytes()
    report = run_pipeline(file_bytes, "sample.docx")

    assert 0 <= report.overall_score <= 100
    assert report.parse_completeness.total_score > 0


def test_pdf_two_column_reading_order_is_not_interleaved():
    """
    The critical regression test: a naive top-to-bottom-only extractor
    would interleave 'Experience' / 'Skills' / 'Senior Engineer...' /
    'Python, SQL, AWS' by y-position. Correct column-aware extraction
    must read the entire left column before the right column.
    """
    pdf_bytes = _build_two_column_pdf_bytes()
    spans = extract_spans(pdf_bytes)
    dims = page_dimensions(pdf_bytes)
    ordered = reconstruct_reading_order(spans, dims)
    lines = spans_to_lines(ordered)

    left_col_end_index = lines.index("Senior Engineer, TechCo")
    right_col_start_index = lines.index("Skills")

    # Everything in the left column must come before anything in the
    # right column - i.e. reading order is NOT simple y-position order.
    assert left_col_end_index < right_col_start_index


def test_bullet_marker_stripping_does_not_break_verb_detection():
    """
    Regression test for the bug found during manual testing: leading
    '-' bullet markers were being treated as the 'first word', causing
    every bullet to fail strong-verb detection.
    """
    bullets = [
        "- Led a team of 5 engineers to launch a new billing platform",
        "- Responsible for roadmap planning across three product lines",
        "- Reduced churn by 15% through onboarding redesign",
    ]
    result = analyze_bullets(bullets)

    assert result.strong_verb_start_count == 2  # "Led", "Reduced"
    assert result.weak_opener_count == 1  # "Responsible for"


def test_unsupported_file_type_raises_value_error():
    import pytest

    with pytest.raises(ValueError):
        run_pipeline(b"not a real file", "resume.txt")
