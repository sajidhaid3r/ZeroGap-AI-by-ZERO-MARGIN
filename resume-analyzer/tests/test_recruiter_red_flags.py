"""
Tests for the recruiter-red-flags round: format/design risk, resume
length, overlapping-employment detection, buzzword-without-evidence
detection, and the expanded 7-category composite score - including a
regression test for a real bug found during this round (buzzwords in
the summary section were being silently skipped).
"""

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent.parent))

import io
from docx import Document

from app.scoring.format_risk import run_format_risk_checks
from app.scoring.length_check import check_resume_length
from app.scoring.overlap_detector import detect_overlaps
from app.scoring.buzzword_detector import detect_buzzwords
from app.scoring.composite_score import CATEGORY_WEIGHTS
from app.pipeline import run_pipeline


# --- Format risk ---

def test_multi_column_flagged():
    result = run_format_risk_checks([], column_count=2)
    col_check = [c for c in result.checks if c.name == "multi_column_layout"][0]
    assert col_check.passed is False


def test_single_column_passes():
    result = run_format_risk_checks([], column_count=1)
    col_check = [c for c in result.checks if c.name == "multi_column_layout"][0]
    assert col_check.passed is True


def test_embedded_images_flagged():
    result = run_format_risk_checks([], image_count=2)
    img_check = [c for c in result.checks if c.name == "embedded_images"][0]
    assert img_check.passed is False


def test_tables_used_flagged():
    result = run_format_risk_checks([], table_line_count=5)
    table_check = [c for c in result.checks if c.name == "tables_used_for_layout"][0]
    assert table_check.passed is False


# --- Length check ---

def test_short_resume_flagged():
    result = check_resume_length(["word " * 30])
    assert result.assessment == "too_short"


def test_ideal_length_resume():
    result = check_resume_length(["word " * 400])
    assert result.assessment == "ideal"


def test_too_long_resume_flagged():
    result = check_resume_length(["word " * 1500])
    assert result.assessment == "too_long"


# --- Overlap detection ---

def test_overlapping_roles_detected():
    ranges = [
        ("Senior Engineer at Acme", "Jan 2022", "Present"),
        ("Freelance Consultant", "Mar 2021", "Dec 2022"),
    ]
    overlaps = detect_overlaps(ranges)
    assert len(overlaps) == 1
    assert overlaps[0].overlap_months > 0


def test_non_overlapping_roles_not_flagged():
    ranges = [
        ("Senior Engineer at Acme", "Jan 2022", "Present"),
        ("Engineer at Beta", "Jun 2018", "Nov 2021"),
    ]
    overlaps = detect_overlaps(ranges)
    assert len(overlaps) == 0


# --- Buzzword detection ---

def test_buzzword_without_evidence_flagged():
    result = detect_buzzwords(["Results-driven team player."])
    assert result.unsupported_count == 2  # "Results-driven" and "team player"


def test_buzzword_with_evidence_not_counted_as_unsupported():
    result = detect_buzzwords(["Led a team of 5 engineers, a results-driven approach that cut costs by 15%"])
    finding = [f for f in result.findings if f.buzzword.lower() == "results-driven"][0]
    assert finding.has_nearby_evidence is True


def test_buzzwords_in_summary_are_scanned_via_pipeline():
    """Regression test for a real bug found during development: the
    pipeline originally only scanned experience bullets for buzzwords,
    silently missing the summary/objective section - one of the most
    buzzword-dense parts of a typical resume."""
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("jane.doe@proton.me")
    doc.add_paragraph("Summary")
    doc.add_paragraph("Results-driven team player with excellent communication skills.")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Product Manager, Acme Corp, Jan 2022 - Present")
    doc.add_paragraph("- Led product launches")
    doc.add_paragraph("Skills")
    doc.add_paragraph("SQL")

    buf = io.BytesIO()
    doc.save(buf)

    report = run_pipeline(buf.getvalue(), "resume.docx")
    buzzwords_found = {f.buzzword.lower() for f in report.buzzwords.findings}
    assert "results-driven" in buzzwords_found
    assert "team player" in buzzwords_found


# --- Composite score with red flags ---

def test_category_weights_still_sum_to_one_with_red_flags_category():
    assert "recruiter_red_flags" in CATEGORY_WEIGHTS
    assert abs(sum(CATEGORY_WEIGHTS.values()) - 1.0) < 1e-6


def test_pipeline_detects_overlap_end_to_end():
    doc = Document()
    doc.add_paragraph("Jane Doe")
    doc.add_paragraph("jane.doe@proton.me")
    doc.add_paragraph("Experience")
    doc.add_paragraph("Senior Engineer, Acme Corp, Jan 2022 - Present")
    doc.add_paragraph("- Built backend systems for the payments team")
    doc.add_paragraph("Freelance Consultant, Self-Employed, Mar 2021 - Dec 2022")
    doc.add_paragraph("- Advised startups on architecture decisions")
    doc.add_paragraph("Skills")
    doc.add_paragraph("Python")

    buf = io.BytesIO()
    doc.save(buf)

    report = run_pipeline(buf.getvalue(), "resume.docx")
    assert len(report.employment_overlaps) == 1
