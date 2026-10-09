"""
DOCX extraction. Word files are already structured XML, so - unlike PDFs -
there's no reading-order ambiguity to solve. The main risks here instead
are tables (used as layout hacks) and headers/footers (where contact info
sometimes hides and many parsers never look).
"""

from dataclasses import dataclass
from typing import List

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


@dataclass
class DocxLine:
    text: str
    is_bold: bool
    is_heading_style: bool
    source: str  # "body" | "table" | "header" | "footer"


class DocxExtractionError(Exception):
    pass


def _paragraph_is_bold(paragraph: Paragraph) -> bool:
    return any(run.bold for run in paragraph.runs if run.text.strip())


def _paragraph_is_heading(paragraph: Paragraph) -> bool:
    style_name = (paragraph.style.name or "").lower()
    return "heading" in style_name or "title" in style_name


def extract_lines(docx_bytes: bytes) -> List[DocxLine]:
    import io

    try:
        doc = Document(io.BytesIO(docx_bytes))
    except Exception as e:
        raise DocxExtractionError(f"Could not open DOCX file: {e}")

    lines: List[DocxLine] = []

    # Body paragraphs, in document order
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        lines.append(
            DocxLine(
                text=text,
                is_bold=_paragraph_is_bold(para),
                is_heading_style=_paragraph_is_heading(para),
                source="body",
            )
        )

    # Tables - flag explicitly. A resume relying on tables for layout is
    # an ATS risk in itself (this feeds the "ATS-friendly design" check
    # downstream), so we tag these lines rather than silently merging them.
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                text = cell.text.strip()
                if text:
                    lines.append(
                        DocxLine(text=text, is_bold=False, is_heading_style=False, source="table")
                    )

    # Headers/footers - contact info hidden here is invisible to many
    # real ATS parsers, so we extract it but will flag it as a risk,
    # not silently treat it as equivalent to body text.
    for section in doc.sections:
        header = section.header
        for para in header.paragraphs:
            text = para.text.strip()
            if text:
                lines.append(
                    DocxLine(text=text, is_bold=False, is_heading_style=False, source="header")
                )
        footer = section.footer
        for para in footer.paragraphs:
            text = para.text.strip()
            if text:
                lines.append(
                    DocxLine(text=text, is_bold=False, is_heading_style=False, source="footer")
                )

    if not lines:
        raise DocxExtractionError("No extractable text found in DOCX file.")

    return lines


def count_embedded_images(docx_bytes: bytes) -> int:
    """Counts inline images embedded in the document - used by the
    format-risk check, since text inside images is invisible to ATS
    parsers."""
    import io

    doc = Document(io.BytesIO(docx_bytes))
    return len(doc.inline_shapes)
